"""
Mapping Sheet Automation - single file Streamlit app

Run:
  pip install -r requirements.txt
  streamlit run app.py

Inputs
  1. Label to group file (.xlsx): parent group ek column me, child groups uske aage ke columns me
     (long format bhi chalega)
  2. CGC file(s) (.xlsx): category_name, group_name, class_name, class_image_gcs_file_path
     (parent aur child dono ki images)

Output
  mapping_sheet.xlsx: Parent group | Parent Class | Child Group | Child Class | Score | Status | Suggestion
"""
import hashlib
import io
import os

import numpy as np
import pandas as pd
import requests
import streamlit as st
from PIL import Image

HIGH_DEFAULT = 0.85   # score >= HIGH -> same product family (auto match)
LOW_DEFAULT = 0.70    # score <  LOW  -> different (others); beech ka score -> REVIEW
CACHE_DIR = "image_cache"


def norm(s):
    return " ".join(str(s).lower().split())


# ------------------------------------------------------------------ label file -> long rows
def read_label_file(file, parent_col="A"):
    raw = pd.read_excel(file, header=None, dtype=str)
    pidx = ord(parent_col.strip().upper()) - ord("A")
    rows = []
    for _, r in raw.iterrows():
        parent = r.iloc[pidx] if pidx < len(r) else None
        if pd.isna(parent) or not str(parent).strip():
            continue
        parent = str(parent).strip()
        if norm(parent) in ("parent groups", "parent group"):
            continue
        for child in r.iloc[pidx + 1:]:
            if pd.notna(child) and str(child).strip():
                if norm(child) in ("child groups", "child group"):
                    continue
                rows.append((parent, str(child).strip()))
    return pd.DataFrame(rows, columns=["parent", "child"])


# ------------------------------------------------------------------ images + embeddings
@st.cache_resource(show_spinner="CLIP model load ho raha hai (pehli baar time lagega)...")
def get_model():
    from transformers import CLIPModel, CLIPProcessor
    return (CLIPModel.from_pretrained("openai/clip-vit-base-patch32"),
            CLIPProcessor.from_pretrained("openai/clip-vit-base-patch32"))


def load_image(url):
    os.makedirs(CACHE_DIR, exist_ok=True)
    fp = os.path.join(CACHE_DIR, hashlib.md5(url.encode()).hexdigest() + ".png")
    if os.path.exists(fp):
        return Image.open(fp).convert("RGB")
    resp = requests.get(url, timeout=30)
    resp.raise_for_status()
    with open(fp, "wb") as f:
        f.write(resp.content)
    return Image.open(io.BytesIO(resp.content)).convert("RGB")


def embed_urls(urls, progress=None, batch=32):
    """Normalized CLIP embeddings (n, 512). Download fail hone par zero row."""
    import torch
    model, proc = get_model()
    out = []
    for i in range(0, len(urls), batch):
        imgs, ok = [], []
        for u in urls[i:i + batch]:
            try:
                imgs.append(load_image(u))
                ok.append(True)
            except Exception:
                ok.append(False)
        emb = np.zeros((len(ok), 512), dtype="float32")
        if imgs:
            with torch.no_grad():
                e = model.get_image_features(**proc(images=imgs, return_tensors="pt"))
            emb[np.array(ok)] = torch.nn.functional.normalize(e, dim=1).numpy()
        out.append(emb)
        if progress:
            done = min(i + batch, len(urls))
            progress.progress(done / len(urls), text=f"Images processed: {done}/{len(urls)}")
    return np.vstack(out) if out else np.zeros((0, 512), dtype="float32")


# ------------------------------------------------------------------ main matching logic
def build_mapping(label_df, cgc, embed_fn, high, low):
    cgc = cgc.copy().reset_index(drop=True)
    cgc["g"] = cgc["group_name"].map(norm)
    cgc["c"] = cgc["class_name"].map(norm)
    cgc["emb_i"] = range(len(cgc))
    E = embed_fn(cgc["class_image_gcs_file_path"].tolist())

    group_idx = cgc.groupby("g")["emb_i"].apply(list).to_dict()
    parent_keys = label_df["parent"].map(norm).unique().tolist()
    all_parent_keys = [p for p in parent_keys if p in group_idx]

    def best_class(child_emb, parent_key):
        sub = cgc[cgc["g"] == parent_key]
        sims = child_emb @ E[sub["emb_i"].values].T
        j = np.unravel_index(sims.argmax(), sims.shape)[1]
        return sub.iloc[j]["class_name"], float(sims.max())

    rows = []
    for _, r in label_df.iterrows():
        parent, child = r["parent"], r["child"]
        pk, ck = norm(parent), norm(child)

        idx = group_idx.get(ck) or cgc.loc[cgc["c"] == ck, "emb_i"].tolist()
        idx = [i for i in idx if E[i].any()]
        if not idx:
            rows.append(["", "", child, child, None, "NO IMAGE",
                         "child ki images CGC file me nahi mili"])
            continue
        ce = E[idx]

        own_name, own_score = None, -1.0
        if pk in group_idx and pk != ck:
            own_name, own_score = best_class(ce, pk)

        other_score, other_parent = -1.0, None
        for k in all_parent_keys:
            if k in (pk, ck):
                continue
            _, sc = best_class(ce, k)
            if sc > other_score:
                other_score, other_parent = sc, k
        other_disp = (cgc.loc[cgc["g"] == other_parent, "group_name"].iloc[0]
                      if other_parent else None)

        if own_score >= high:
            rows.append([parent, own_name, child, child, round(own_score, 3), "AUTO-MATCH", ""])
        elif own_score >= low:
            rows.append([parent, own_name, child, child, round(own_score, 3), "REVIEW",
                         "score beech me hai"])
        else:
            status, sug = "AUTO-OTHERS", ""
            if other_score >= high:
                status = "REVIEW"
                sug = f"shayad '{other_disp}' ({other_score:.2f})"
            rows.append(["others", "others", child, child,
                         round(own_score, 3) if own_score > -1 else None, status, sug])

    return pd.DataFrame(rows, columns=["Parent group", "Parent Class", "Child Group",
                                       "Child Class", "Score", "Status", "Suggestion"])


# ------------------------------------------------------------------ excel output
def to_excel_bytes(df):
    from openpyxl.styles import PatternFill
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False)
        ws = w.sheets["Sheet1"]
        fills = {"REVIEW": "FFF2CC", "NO IMAGE": "F4CCCC", "AUTO-OTHERS": "E7E6E6"}
        for row in ws.iter_rows(min_row=2):
            f = fills.get(row[5].value)
            if f:
                for c in row:
                    c.fill = PatternFill("solid", fgColor=f)
        for col, wd in zip("ABCDEFG", (38, 38, 38, 38, 8, 14, 45)):
            ws.column_dimensions[col].width = wd
    return buf.getvalue()


# ================================================================== UI
st.set_page_config(page_title="Mapping Sheet Automation", layout="wide")
st.title("Mapping Sheet Automation")
st.caption("Label to group file + CGC file (class images) se image-based mapping sheet")

with st.sidebar:
    st.header("Settings")
    parent_col = st.text_input("Parent group column (label file)", value="A",
                               help="Jis column me Parent group likha hai uska letter. "
                                    "Child uske aage ke columns me hote hain.")
    high = st.slider("HIGH: is score ya upar = auto match", 0.50, 1.00, HIGH_DEFAULT, 0.01)
    low = st.slider("LOW: is score se neeche = others", 0.30, 1.00, LOW_DEFAULT, 0.01)
    if low > high:
        st.error("LOW, HIGH se chhota hona chahiye")
    st.markdown("Beech ke score wali rows **REVIEW** me aati hain.")

c1, c2 = st.columns(2)
with c1:
    label_file = st.file_uploader("1. Label to group file (.xlsx)", type=["xlsx"])
with c2:
    cgc_files = st.file_uploader(
        "2. CGC file(s) (.xlsx): parent aur child dono ki images",
        type=["xlsx"], accept_multiple_files=True,
        help="Columns: category_name, group_name, class_name, class_image_gcs_file_path. "
             "Parent (ABC) aur child ki images alag files me hon to dono upload kar dein.")

label_df = None
if label_file:
    label_df = read_label_file(label_file, parent_col)
    with st.expander(f"Label file preview ({len(label_df)} parent-child rows)"):
        st.dataframe(label_df, use_container_width=True, hide_index=True)

cgc = None
if cgc_files:
    cgc = pd.concat([pd.read_excel(f) for f in cgc_files], ignore_index=True)
    need = {"group_name", "class_name", "class_image_gcs_file_path"}
    if not need.issubset(cgc.columns):
        st.error(f"CGC file me ye columns chahiye: {sorted(need)}")
        cgc = None
    else:
        cgc = cgc.dropna(subset=["class_image_gcs_file_path"]).drop_duplicates(
            subset=["group_name", "class_name", "class_image_gcs_file_path"])
        st.caption(f"CGC: {len(cgc)} images, {cgc['group_name'].nunique()} groups")

ready = label_df is not None and cgc is not None and low <= high
if st.button("Mapping sheet banao", type="primary", disabled=not ready):
    progress = st.progress(0.0, text="Images download + embeddings...")
    try:
        result = build_mapping(label_df, cgc,
                               embed_fn=lambda urls: embed_urls(urls, progress),
                               high=high, low=low)
        st.session_state["result"] = result
    except Exception as e:
        st.error(f"Error: {e}")
    finally:
        progress.empty()

if "result" in st.session_state:
    result = st.session_state["result"]
    st.subheader("Result")

    counts = result["Status"].value_counts()
    for col, name in zip(st.columns(4), ["AUTO-MATCH", "AUTO-OTHERS", "REVIEW", "NO IMAGE"]):
        col.metric(name, int(counts.get(name, 0)))

    statuses = sorted(result["Status"].unique())
    show = st.multiselect("Status filter", statuses, default=statuses)
    st.caption("Table editable hai: REVIEW rows me Parent group / Parent Class khud theek kar sakte hain.")
    edited = st.data_editor(result[result["Status"].isin(show)], use_container_width=True,
                            hide_index=True, num_rows="fixed", key="editor")

    final = result.copy()
    final.loc[edited.index, edited.columns] = edited

    st.download_button("Excel download karo (mapping_sheet.xlsx)", to_excel_bytes(final),
                       file_name="mapping_sheet.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
